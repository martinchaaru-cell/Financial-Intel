"""Extended survey fields in the cross-company figures, the Intelligence Report rows and the Excel exports."""
import io

import openpyxl

import market_exports
import survey_aggregate as sa
import survey_extended as se
from app import Financials
from models import Company
from models_survey import SurveyCompanyData


def _co(db, name, cur='KES', unit='millions', pay='thousands', **vals):
    c = Company(name=name, ticker=name[:3].upper(), sector='Banking')
    db.session.add(c)
    db.session.commit()
    db.session.add(SurveyCompanyData(company_id=c.id, fiscal_year='FY2025', currency=cur, unit=unit, director_figures_unit=pay,
                                     turnover=1000, net_profit=100, **vals))
    db.session.commit()
    return c


def test_money_is_converted_to_whole_shillings_and_averaged(client, db):
    _co(db, 'Alpha Bank', audit_fee=50, kmp_short_term_benefits=200_000, employees_count=1000, staff_costs=300,
        auditor_name='KPMG Kenya', audit_opinion='Unmodified', going_concern_uncertainty='No', chair_ceo_separate='Yes')
    _co(db, 'Beta Bank', unit='thousands', pay='units', audit_fee=150_000, kmp_short_term_benefits=400_000_000, employees_count=2000,
        staff_costs=500_000, auditor_name='KPMG Kenya', audit_opinion='Qualified', going_concern_uncertainty='Yes', chair_ceo_separate='No')
    ov = sa.build_survey_overview('FY2025')
    e = ov['extended']
    assert e['audit']['audit_fee']['average'] == 100_000_000           # 50m and 150m
    assert e['executive_pay']['kmp_total']['average'] == 300_000_000   # 200m and 400m (salaries only)
    assert e['workforce']['staff_cost_per_employee']['average'] == 275_000   # 300k and 250k per head
    assert e['audit']['auditor_share'][0] == {**e['audit']['auditor_share'][0], 'label': 'KPMG Kenya', 'companies': 2}
    assert e['audit']['opinions']['Qualified']['company_names'] == ['Beta Bank']
    assert e['audit']['going_concern']['yes_companies'] == ['Beta Bank']
    assert e['board_effectiveness']['chair_ceo_separate']['pct_yes'] == 50.0
    rows = {r['name']: r for r in sa.build_company_rows('FY2025')['rows']}
    assert rows['Alpha Bank']['audit_fee'] == 50_000_000 and rows['Beta Bank']['employees_count'] == 2000


def test_one_company_is_not_an_average(client, db):
    _co(db, 'Solo', audit_fee=10)
    m = sa.build_survey_overview('FY2025')['extended']['audit']['audit_fee']
    assert m['average'] is None and m['company_count'] == 1


def test_other_currency_is_converted(client, db):
    _co(db, 'Rwanda Co', cur='RWF', unit='millions', audit_fee=100)
    _co(db, 'Kenya Co', audit_fee=100)
    _co(db, 'Kenya Two', audit_fee=100)          # KES is the most common currency, so the survey is in KES
    rows = {r['name']: r for r in sa.build_company_rows('FY2025')['rows']}
    assert rows['Kenya Co']['audit_fee'] == 100_000_000
    assert 0 < rows['Rwanda Co']['audit_fee'] < 100_000_000       # RWF is worth less than KES


def test_exports_use_survey_figures_in_full(client, db):
    c = _co(db, 'Gamma', earnings_per_share=2.5, employees_count=300)
    db.session.add(Financials(company_id=c.id, period='FY2025', revenue=0, net_income=0, total_assets=10, total_liabilities=6, total_equity=4))
    db.session.commit()
    ws = openpyxl.load_workbook(market_exports.build_portfolio(Company, Financials))['Portfolio']
    head, row = [c_.value for c_ in ws[1]], [c_.value for c_ in ws[2]]
    got = dict(zip(head, row))
    assert got['Turnover (KES)'] == 1_000_000_000 and got['Net margin %'] == 10.0
    assert got['ROE %'] is None and got['Debt / equity'] == 1.5          # no profit on the statement -> no ROE
    wb = openpyxl.load_workbook(market_exports.build_comparison(Company, Financials))
    assert 'Gamma' in wb.sheetnames and 'Notes' in wb.sheetnames
