"""Extraction report: classification of filled / low-confidence / missing fields and the downloads."""
import io
import json

import extraction_report as er


def _seed(db):
    from models import Company
    from models_survey import SurveyCompanyData
    a = Company(name='Alpha Ltd', ticker='ALP')
    b = Company(name='Beta Plc', ticker='BET')
    db.session.add_all([a, b])
    db.session.commit()
    row = SurveyCompanyData(company_id=a.id, fiscal_year='FY2026', source_filename='alpha.pdf',
                            board_size=9, directors_female=3, chairperson_annual_retainer=500000.0, ceo_monthly_salary=900000.0)
    row.field_sources = {'board_size': 'from the board table (page 4)', 'directors_female': 'from the document scan (page 5, stat block): 3 women',
                         'chairperson_annual_retainer': 'entered manually: form upload', 'ceo_monthly_salary': 'from the document scan (page 80)'}
    row.field_confidence = {'board_size': 0.9, 'directors_female': 0.5, 'chairperson_annual_retainer': 1.0}
    db.session.add(row)
    db.session.add(er.ExtractionTrace(
        company_id=a.id, fiscal_year='FY2026', filename='alpha.pdf', kind='import',
        doc_json=json.dumps({'pages': 100, 'text_pages_pct': 98, 'looks_like_spreads': True, 'statement_unit': 'thousands'}),
        topics_json=json.dumps({'market_cap': {'pages': [7], 'snippet': 'market capitalisation of US$1 billion'}}),
        decisions_json=json.dumps({'rejections': [{'field': 'neds_kenyan_count', 'reason': 'nationality block total differs from board size',
                                                    'value': 4, 'page': 12, 'text': ''}], 'decisions': []})))
    db.session.commit()
    return a.id, b.id


def test_statuses_and_reasons(app, db):
    a, b = _seed(db)
    rep = er.build_report(company_id=a, fiscal_year='FY2026')
    f = rep['files'][0]
    items = {i['field']: i for i in f['fields']}
    assert items['board_size']['status'] == 'filled' and items['board_size']['reason_code'] == 'filled'
    assert items['directors_female']['status'] == 'low_confidence'
    assert items['chairperson_annual_retainer']['reason_code'] == 'filled_manual'
    assert items['ceo_monthly_salary']['status'] == 'filled'          # no confidence stored -> complete, same as data_quality
    # topic present in the text but nothing parsed, and the PDF is two-page spreads
    mc = items['market_cap']
    assert mc['status'] == 'missing' and mc['reason_code'] == 'topic_found' and 'spreads' in mc['reason'] and 'page(s) 7' in mc['reason']
    # not mentioned anywhere in the text
    assert items['ceo_monthly_gratuity']['reason_code'] == 'not_in_text'
    # a reader found a value and dropped it
    k = items['neds_kenyan_count']
    assert k['reason_code'] == 'rejected' and 'page 12' in k['reason']
    assert f['counts']['total'] == 40 and f['counts']['filled'] + f['counts']['low_confidence'] + f['counts']['missing'] == 40


def test_agrees_with_data_quality_score(app, db):
    import data_quality
    from models_survey import SurveyCompanyData
    a, _ = _seed(db)
    row = SurveyCompanyData.query.filter_by(company_id=a).first()
    q = data_quality.company_quality(row)
    f = er.build_report(company_id=a)['files'][0]
    assert f['counts']['filled'] == q['complete'] and f['counts']['low_confidence'] == q['needs_review'] and f['counts']['missing'] == q['missing']


def test_no_trace_and_no_row(app, db):
    a, b = _seed(db)
    from models_survey import SurveyCompanyData
    db.session.add(SurveyCompanyData(company_id=b, fiscal_year='FY2025', board_size=5))
    db.session.commit()
    rep = er.build_report()
    beta = next(f for f in rep['files'] if f['company'] == 'Beta Plc')
    assert beta['trace'] is None
    assert next(i for i in beta['fields'] if i['field'] == 'market_cap')['reason_code'] == 'no_trace'
    # Alpha has no FY2025 row and Beta has no FY2026 row
    missing = {(n['company'], n['fiscal_year']) for n in rep['no_survey_row']}
    assert ('Alpha Ltd', 'FY2025') in missing and ('Beta Plc', 'FY2026') in missing
    assert rep['rollups']['by_reason']['no_trace'] > 0


def test_downloads_and_routes(app, client, db):
    from app import User
    _seed(db)
    u = User(name='adm', email='adm@x.com', role='admin')
    u.set_password('secret1')
    db.session.add(u); db.session.commit()
    assert client.get('/api/extraction-report').status_code == 403            # guests are refused
    assert client.post('/api/login', json={'email': 'adm@x.com', 'password': 'secret1'}).status_code == 200
    assert client.get('/extraction-report').status_code == 200
    opts = client.get('/api/extraction-report/options').get_json()
    assert opts['years'] == ['FY2026'] and len(opts['companies']) == 2
    assert len(client.get('/api/extraction-report?fiscal_year=FY2026').get_json()['files']) == 1
    csv_r = client.get('/api/extraction-report/download?format=csv')
    assert csv_r.status_code == 200 and b'Company' in csv_r.data and b'topic_found' in csv_r.data
    x = client.get('/api/extraction-report/download?format=xlsx')
    assert x.status_code == 200 and x.data[:2] == b'PK'
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(x.data))
    assert {'Summary', 'Fields', 'Weakest fields', 'Failure reasons', 'Upload problems'} <= set(wb.sheetnames)
    assert wb['Fields'].max_row > 40
    assert client.post('/api/extraction-report/diagnose', data={}).status_code == 400


def test_dynamic_reader_records_rejections():
    from dynamic_extract import _Result
    r = _Result()
    r.offer('board_size', 9, 4, 'nine directors', 0.9, 'stated')
    r.offer('board_size', 8, 5, 'eight directors', 0.5, 'roster')
    assert r.fields['board_size']['value'] == 9
    assert r.rejections and r.rejections[0]['field'] == 'board_size' and 'lost' in r.rejections[0]['reason']


def _tiny_pdf():
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    lines = ['Board of Directors', 'The Board comprises nine directors, of whom three are women.',
             'Directors sitting allowance is paid per meeting attended.', 'Market capitalisation was KES 12 billion.']
    for i, t in enumerate(lines):
        c.drawString(60, 780 - i * 20, t)
    c.save()
    return buf.getvalue()


def test_import_hook_and_diagnose_endpoint(app, client, db):
    """Same calls the upload flow makes: the dynamic reader writes + stashes decisions, record_import stores the trace."""
    import app as appmod
    from models import Company
    from models_survey import SurveyCompanyData
    c = Company(name='Gamma', ticker='GAM')
    db.session.add(c); db.session.commit()
    db.session.add(SurveyCompanyData(company_id=c.id, fiscal_year='FY2026'))
    db.session.commit()
    pdf = _tiny_pdf()
    appmod._apply_dynamic_survey(c.id, 'FY2026', None, pdf, 'gamma.pdf')
    assert (c.id, 'FY2026') in er._PENDING
    assert er.record_import(c.id, 'FY2026', 'gamma.pdf', None, pdf)
    assert (c.id, 'FY2026') not in er._PENDING
    f = er.build_report(company_id=c.id)['files'][0]
    items = {i['field']: i for i in f['fields']}
    assert f['trace']['doc']['pages'] == 1
    assert items['market_cap']['reason_code'] == 'topic_found'          # the page mentions it, no reader parsed it
    assert items['ceo_monthly_gratuity']['reason_code'] == 'not_in_text'

    from app import User
    u = User(name='adm', email='a2@x.com', role='admin'); u.set_password('secret1')
    db.session.add(u); db.session.commit()
    client.post('/api/login', json={'email': 'a2@x.com', 'password': 'secret1'})
    r = client.post('/api/extraction-report/diagnose', data={'company_id': str(c.id), 'fiscal_year': 'FY2026',
                                                               'file': (io.BytesIO(pdf), 'gamma.pdf')},
                    content_type='multipart/form-data')
    assert r.status_code == 200 and r.get_json()['ok'] and r.get_json()['report']['files'][0]['trace']['kind'] == 'diagnose'


def test_tables_parity_pdf_and_zip(app, client, db):
    import zipfile
    from app import User
    from models_survey import SurveyCompanyData, SurveyDirector
    a, b = _seed(db)
    db.session.add(SurveyDirector(company_id=a, fiscal_year='FY2026', director_name='Jane Doe', order_index=1))
    db.session.add(SurveyCompanyData(company_id=b, fiscal_year='FY2025', board_size=5))
    db.session.commit()
    f = next(x for x in er.build_report(company_id=a, fiscal_year='FY2026')['files'])
    assert any(t['section'] == "Directors' register" and t['item'] == 'Jane Doe' for t in f['tables'])
    assert any(t['item'] == 'Table is empty' for t in f['tables'])            # committees / pay rows / benefits not on file
    assert f['parity']['agree'] is True                                       # report and Manual Entry form list the same missing fields
    buf, name = er.company_pdf(a, 'FY2026')
    assert buf.read(4) == b'%PDF' and name == 'Text_Extraction_Report_Alpha_Ltd_FY2026.pdf'
    assert er.company_pdf(a, 'FY2099') == (None, None)
    zbuf, zname = er.build_zip()
    names = zipfile.ZipFile(zbuf).namelist()
    assert '00_Extraction_Coverage_Summary.pdf' in names and '00_Extraction_Coverage_Summary.csv' in names
    assert sum(1 for n in names if n.startswith(('01_', '02_'))) == 2 and zname.startswith('Text_Extraction_Reports_All_years_')
    z2 = zipfile.ZipFile(er.build_zip('FY2025')[0])
    assert len([n for n in z2.namelist() if 'Beta' in n]) == 1 and not [n for n in z2.namelist() if 'Alpha' in n]
    u = User(name='adm', email='adm@x.com', role='admin'); u.set_password('secret1'); db.session.add(u); db.session.commit()
    assert client.get(f'/api/extraction-report/company/{a}.pdf').status_code == 403
    client.post('/api/login', json={'email': 'adm@x.com', 'password': 'secret1'})
    r = client.get(f'/api/extraction-report/company/{a}.pdf?fiscal_year=FY2026')
    assert r.status_code == 200 and r.data[:4] == b'%PDF'
    assert client.get(f'/api/extraction-report/company/{a}.pdf?fiscal_year=FY2099').status_code == 404
    r = client.get('/api/extraction-report/zip?fiscal_year=FY2026')
    assert r.status_code == 200 and r.data[:2] == b'PK'


def test_bulk_diagnose_matches_file_names_and_stores_traces(app, client, db):
    import zipfile
    from app import User
    from models import Company
    from models_survey import SurveyCompanyData
    k = Company(name='KCB Group', ticker='KCB'); s = Company(name='Safaricom', ticker='SCOM'); e = Company(name='East African Breweries', ticker='EABL')
    db.session.add_all([k, s, e]); db.session.commit()
    db.session.add(SurveyCompanyData(company_id=k.id, fiscal_year='FY2025', board_size=9))
    db.session.add(SurveyCompanyData(company_id=s.id, fiscal_year='FY2026', board_size=9))     # year ends the following calendar year
    db.session.commit()
    assert er.match_filename('ke-kcb-2025-ar-00.pdf')[:2] == (k, 'FY2025')
    co, fy, note = er.match_filename('ke-scom-2025-ar-00.pdf')
    assert co == s and fy == 'FY2026' and 'used FY2026' in note
    assert er.match_filename('ke-eabl-2025-ar-00.pdf')[0] == e
    assert er.match_filename('random.pdf')[0] is None
    pdf = _tiny_pdf()
    res = er.diagnose_many([('ke-kcb-2025-ar-00.pdf', pdf), ('ke-zzz-2025-ar-00.pdf', pdf), ('ke-eabl-2025-ar-00.pdf', b'not a pdf')])
    assert [r['status'] for r in res] == ['ok', 'unmatched', 'not_pdf']
    f = er.build_report(company_id=k.id)['files'][0]
    assert f['trace']['kind'] == 'diagnose' and f['trace']['doc']['pages'] == 1
    assert next(i for i in f['fields'] if i['field'] == 'market_cap')['reason_code'] == 'topic_found'    # cause now known, not 'no_trace'
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, 'w') as z:
        z.writestr('reports/ke-scom-2025-ar-00.pdf', pdf)
    assert er.diagnose_many([('all.zip', buf.getvalue())])[0]['fiscal_year'] == 'FY2026'
    u = User(name='adm', email='adm@x.com', role='admin'); u.set_password('secret1'); db.session.add(u); db.session.commit()
    assert client.post('/api/extraction-report/diagnose-bulk', data={}).status_code == 403
    client.post('/api/login', json={'email': 'adm@x.com', 'password': 'secret1'})
    assert client.post('/api/extraction-report/diagnose-bulk', data={}).status_code == 400
    r = client.post('/api/extraction-report/diagnose-bulk', data={'files': (io.BytesIO(pdf), 'ke-kcb-2025-ar-00.pdf')},
                    content_type='multipart/form-data')
    assert r.status_code == 200 and r.get_json()['results'][0]['status'] == 'ok'